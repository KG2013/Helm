#!/usr/bin/env python3
"""Bounded, one-request-at-a-time document worker.

The worker has no Agent loop, provider credentials, approval authority, or
Runtime state. It receives a JSON object, performs one bounded operation under
HELM_WORKSPACE_ROOT, and emits one JSON response with a redacted artifact
receipt. Optional DOCX/XLSX/PDF dependencies are detected per operation.
"""
from __future__ import annotations

import hashlib
import json
import os
import shutil
import subprocess
import sys
import tempfile
import time
import zipfile
import xml.etree.ElementTree as ET
from html import escape
from pathlib import Path
from typing import Any

MAX_BYTES = 1_000_000
VERSION = "0.3.0"
MAX_XLSX_CELLS = 1_000
MAX_RENDER_SECONDS = 15
MAX_OCR_SECONDS = 20
MIN_OCR_CONFIDENCE = 0.5


def response(request_id: str, *, ok: bool, result: Any = None, error: str | None = None, receipt: dict[str, Any] | None = None) -> dict[str, Any]:
    payload: dict[str, Any] = {"id": request_id, "ok": ok}
    if result is not None:
        payload["result"] = result
    if not ok:
        payload["error"] = error or "worker_error"
    if receipt is not None:
        payload["receipt"] = receipt
    return payload


def workspace_root() -> Path:
    return Path(os.environ.get("HELM_WORKSPACE_ROOT", os.getcwd())).expanduser().resolve()


def safe_path(raw_path: Any, *, must_exist: bool = False) -> Path:
    if not isinstance(raw_path, str) or not raw_path or "\x00" in raw_path:
        raise ValueError("path_required")
    root = workspace_root()
    candidate = (root / raw_path).resolve(strict=False)
    try:
        candidate.relative_to(root)
    except ValueError as exc:
        raise ValueError("path_outside_workspace") from exc
    if candidate.exists() and candidate.is_symlink():
        raise ValueError("symlink_rejected")
    if must_exist and not candidate.is_file():
        raise ValueError("file_not_found")
    if candidate.exists() and candidate.stat().st_size > MAX_BYTES:
        raise ValueError("file_exceeds_bound")
    return candidate


def artifact(path: Path, *, operation: str, source_run_id: Any = None, limitations: list[str] | None = None) -> dict[str, Any]:
    relative_path = path.relative_to(workspace_root()).as_posix()
    # Re-canonicalize immediately before hashing so a swapped symlink cannot
    # turn the receipt into an artifact outside the authorized workspace.
    verified_path = safe_path(relative_path, must_exist=True)
    data = verified_path.read_bytes()
    if len(data) > MAX_BYTES:
        raise ValueError("artifact_exceeds_bound")
    return {
        "type": verified_path.suffix.lower().lstrip(".") or "file",
        "path": relative_path,
        "hash": hashlib.sha256(data).hexdigest(),
        "bytes": len(data),
        "operation": operation,
        "sourceRunId": source_run_id,
        "workerVersion": VERSION,
        "limitations": limitations or [],
    }


def receipt(path: Path, *, operation: str, request: dict[str, Any], side_effect: str = "known", limitations: list[str] | None = None, checks: dict[str, str] | None = None, target: dict[str, Any] | None = None) -> dict[str, Any]:
    value: dict[str, Any] = {"worker": "document-worker", "workerVersion": VERSION, "operation": operation, "sideEffect": side_effect, "artifact": artifact(path, operation=operation, source_run_id=request.get("runId"), limitations=limitations)}
    if checks is not None:
        value["checks"] = checks
    if target is not None:
        value["target"] = target
    return value


def handle(request: dict[str, Any]) -> dict[str, Any]:
    request_id = str(request.get("id", "unknown"))
    operation = request.get("operation")
    if operation == "health":
        return response(request_id, ok=True, result={"worker": "document-worker", "version": VERSION})
    try:
        if operation == "inspect":
            path = safe_path(request.get("path"), must_exist=True)
            return response(request_id, ok=True, result={"path": path.relative_to(workspace_root()).as_posix(), "suffix": path.suffix.lower(), "size": path.stat().st_size}, receipt=receipt(path, operation="inspect", request=request, side_effect="none"))
        if operation == "docx_create":
            path = safe_path(request.get("path"))
            paragraphs = request.get("paragraphs", [])
            if not isinstance(paragraphs, list) or not all(isinstance(item, str) for item in paragraphs) or len(paragraphs) > 200:
                return response(request_id, ok=False, error="paragraphs_invalid")
            path.parent.mkdir(parents=True, exist_ok=True)
            write_minimal_docx(path, [str(item)[:4_000] for item in paragraphs])
            checks, limitations, rendering = verify_docx(path, [str(item)[:4_000] for item in paragraphs])
            docx_receipt = receipt(path, operation="docx_create", request=request, limitations=limitations, checks=checks)
            docx_receipt["rendering"] = rendering
            if any(status in {"unknown", "conflict"} for status in checks.values()):
                docx_receipt["verification"] = "unknown"
            if checks.get("structure") != "passed" or checks.get("content") != "passed":
                return response(request_id, ok=False, error="docx_verification_unknown", receipt=docx_receipt)
            return response(request_id, ok=True, result={"paragraphs": len(paragraphs), "checks": checks}, receipt=docx_receipt)
        if operation in {"xlsx_read_range", "xlsx_write_range"}:
            return handle_xlsx(request, request_id, operation)
        if operation == "pdf_extract":
            return handle_pdf(request, request_id)
        return response(request_id, ok=False, error="unsupported_operation")
    except ImportError as exc:
        return response(request_id, ok=False, error=f"dependency_missing:{exc.name or 'optional'}")
    except (OSError, ValueError, KeyError) as exc:
        return response(request_id, ok=False, error=str(exc))
    except Exception:
        return response(request_id, ok=False, error="worker_operation_failed")


def handle_xlsx(request: dict[str, Any], request_id: str, operation: str) -> dict[str, Any]:
    from openpyxl import load_workbook

    path = safe_path(request.get("path"), must_exist=True)
    sheet = request.get("sheet")
    selector = xlsx_selector(request)
    if not isinstance(sheet, str) or selector is None:
        return response(request_id, ok=False, error="sheet_and_cell_or_range_required")
    if operation == "xlsx_read_range":
        workbook = load_workbook(path, read_only=True, data_only=False)
        if sheet not in workbook.sheetnames:
            return response(request_id, ok=False, error="sheet_not_found")
        value = xlsx_values(workbook[sheet], selector)
        target = {"sheet": sheet, "selector": selector, "cells": xlsx_cells(selector)}
        checks = {"target": "passed"}
        return response(request_id, ok=True, result={"sheet": sheet, "selector": selector, "value": value, "checks": checks}, receipt=receipt(path, operation=operation, request=request, side_effect="none", checks=checks, target=target))
    workbook = load_workbook(path, data_only=False)
    if sheet not in workbook.sheetnames:
        return response(request_id, ok=False, error="sheet_not_found")
    worksheet = workbook[sheet]
    before = xlsx_snapshot(workbook)
    before_target = xlsx_values(worksheet, selector)
    try:
        expected = xlsx_write_values(worksheet, selector, request)
    except ValueError as exc:
        return response(request_id, ok=False, error=str(exc))
    workbook.save(path)
    after_workbook = load_workbook(path, data_only=False)
    after = xlsx_snapshot(after_workbook)
    actual = xlsx_values(after_workbook[sheet], selector)
    target_passed = actual == expected
    scope_passed = before is not None and after is not None and changed_outside_target(before, after, sheet, selector) == []
    checks = {"target": "passed" if target_passed else "conflict", "scope": "passed" if scope_passed else "unknown" if before is None or after is None else "conflict"}
    target = {"sheet": sheet, "selector": selector, "cells": xlsx_cells(selector), "before": before_target, "after": actual}
    result = {"sheet": sheet, "selector": selector, "value": actual, "checks": checks}
    limitations = [] if scope_passed else ["Workbook scope could not be fully compared; unauthorized changes remain unverified."]
    return response(request_id, ok=True, result=result, receipt=receipt(path, operation=operation, request=request, checks=checks, target=target, limitations=limitations))


def handle_pdf(request: dict[str, Any], request_id: str) -> dict[str, Any]:
    from pypdf import PdfReader

    path = safe_path(request.get("path"), must_exist=True)
    reader = PdfReader(str(path))
    pages: list[dict[str, Any]] = []
    for index, page in enumerate(reader.pages[:100]):
        text = (page.extract_text() or "").strip()
        pages.append({"page": index + 1, "text": text[:10_000], "hasTextLayer": bool(text), "extraction": "text-layer" if text else "unknown", "confidence": 1.0 if text else None, "source": {"path": path.relative_to(workspace_root()).as_posix(), "page": index + 1, "extraction": "text-layer" if text else "unknown"}})
    missing_pages = [page["page"] for page in pages if not page["hasTextLayer"]]
    ocr_limitations: list[str] = []
    ocr_metadata: dict[str, Any] = {"requestedPages": missing_pages, "status": "not_needed" if not missing_pages else "unknown", "limitations": []}
    if missing_pages:
        ocr_pages, ocr_metadata, ocr_limitations = ocr_missing_pages(path, missing_pages)
        ocr_metadata["limitations"] = ocr_limitations
        if ocr_pages is not None:
            for page in pages:
                replacement = ocr_pages.get(page["page"])
                if replacement is not None:
                    page.update(replacement)
    complete = bool(pages) and all(bool(page.get("text")) and (page.get("hasTextLayer") or (page.get("extraction") == "ocr" and isinstance(page.get("confidence"), (int, float)) and page["confidence"] >= MIN_OCR_CONFIDENCE)) for page in pages)
    sources_complete = bool(pages) and all(isinstance(page.get("source"), dict) and isinstance(page["source"].get("page"), int) for page in pages)
    checks = {"coverage": "passed" if complete else "unknown", "sources": "passed" if sources_complete else "unknown"}
    limitations = ocr_limitations if not complete else []
    pdf_receipt = receipt(path, operation="pdf_extract", request=request, side_effect="none", limitations=limitations, checks=checks, target={"pages": [page["page"] for page in pages]})
    pdf_receipt["ocr"] = ocr_metadata
    result = {"pages": pages, "extraction": "text-layer" if not missing_pages else "text-layer+ocr", "checks": checks, "ocr": ocr_metadata}
    if not complete:
        pdf_receipt["verification"] = "unknown"
        return response(request_id, ok=False, result=result, error="ocr_unavailable" if missing_pages and ocr_metadata.get("status") != "passed" else "unknown_text_layer", receipt=pdf_receipt)
    return response(request_id, ok=True, result=result, receipt=pdf_receipt)


def verify_docx(path: Path, paragraphs: list[str]) -> tuple[dict[str, str], list[str], dict[str, Any]]:
    """Validate package/XML/content and, when available, bounded PDF rendering."""
    limitations: list[str] = []
    structure = "passed"
    content = "passed"
    try:
        with zipfile.ZipFile(path) as archive:
            required = {"[Content_Types].xml", "_rels/.rels", "word/document.xml"}
            if not required.issubset(set(archive.namelist())) or archive.testzip() is not None:
                structure = "conflict"
            else:
                root = ET.fromstring(archive.read("word/document.xml"))
                namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
                found = ["".join(node.itertext()) for node in root.iter(f"{namespace}p")]
                if found != paragraphs:
                    content = "conflict"
    except (OSError, ET.ParseError, zipfile.BadZipFile, KeyError):
        structure = "conflict"
    rendering, rendering_evidence, rendering_limitation = render_docx(path)
    if rendering_limitation:
        limitations.append(rendering_limitation)
    return {"structure": structure, "content": content, "rendering": rendering}, limitations, rendering_evidence


def render_docx(path: Path) -> tuple[str, dict[str, Any], str | None]:
    renderer = shutil.which("soffice") or shutil.which("libreoffice")
    if not renderer:
        return "unknown", {"status": "unavailable"}, "No office renderer is available; visual rendering remains unverified."
    with tempfile.TemporaryDirectory(prefix="helm-docx-render-") as temporary:
        output_dir = Path(temporary)
        command = [renderer, "--headless", "--nologo", "--nodefault", "--nofirststartwizard", "--nolockcheck", "--convert-to", "pdf", "--outdir", str(output_dir), str(path)]
        try:
            completed = subprocess.run(command, cwd=temporary, env=command_environment(), stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=MAX_RENDER_SECONDS, check=False)
        except (OSError, subprocess.TimeoutExpired):
            return "unknown", {"status": "failed", "renderer": Path(renderer).name}, "DOCX rendering command failed or exceeded its time bound."
        rendered = output_dir / f"{path.stem}.pdf"
        try:
            rendered_bytes = rendered.stat().st_size if rendered.is_file() else None
        except OSError:
            rendered_bytes = None
        if completed.returncode != 0 or rendered_bytes is None or rendered_bytes > MAX_BYTES:
            return "unknown", {"status": "failed", "renderer": Path(renderer).name}, "DOCX renderer did not produce a bounded PDF artifact."
        try:
            from pypdf import PdfReader
            pages = len(PdfReader(str(rendered)).pages)
        except Exception:
            return "unknown", {"status": "failed", "renderer": Path(renderer).name}, "Rendered PDF could not be opened for a page-count check."
        if pages < 1:
            return "unknown", {"status": "failed", "renderer": Path(renderer).name}, "Rendered PDF has no pages."
        return "passed", {"status": "passed", "renderer": Path(renderer).name, "pdfPages": pages}, "Rendering evidence proves bounded PDF conversion and opening; pixel-level visual comparison is not performed."


def ocr_missing_pages(path: Path, page_numbers: list[int]) -> tuple[dict[int, dict[str, Any]] | None, dict[str, Any], list[str]]:
    pdftoppm = shutil.which("pdftoppm")
    tesseract = shutil.which("tesseract")
    if not pdftoppm or not tesseract:
        missing = [name for name, value in (("pdftoppm", pdftoppm), ("tesseract", tesseract)) if not value]
        limitation = f"OCR is unavailable; missing command(s): {', '.join(missing)}."
        return None, {"status": "unavailable", "missing": missing, "requestedPages": page_numbers, "limitations": [limitation]}, [limitation]
    deadline = time.monotonic() + MAX_OCR_SECONDS
    pages: dict[int, dict[str, Any]] = {}
    limitations: list[str] = []
    with tempfile.TemporaryDirectory(prefix="helm-pdf-ocr-") as temporary:
        output_dir = Path(temporary)
        for page_number in page_numbers:
            if time.monotonic() >= deadline:
                limitations.append("OCR exceeded the total time bound before all pages were processed.")
                break
            image_prefix = output_dir / f"page-{page_number}"
            remaining = deadline - time.monotonic()
            if remaining <= 0:
                limitations.append("OCR exceeded the total time bound before all pages were processed.")
                break
            try:
                rendered = subprocess.run([pdftoppm, "-f", str(page_number), "-l", str(page_number), "-singlefile", "-png", "-r", "150", str(path), str(image_prefix)], cwd=temporary, env=command_environment(), stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=min(10, remaining), check=False)
            except (OSError, subprocess.TimeoutExpired):
                limitations.append(f"OCR page {page_number} rendering failed or timed out.")
                continue
            image = image_prefix.with_suffix(".png")
            if rendered.returncode != 0 or not image.is_file() or image.stat().st_size > MAX_BYTES:
                limitations.append(f"OCR page {page_number} image was unavailable or exceeded its bound.")
                continue
            remaining = deadline - time.monotonic()
            if remaining <= 0:
                limitations.append("OCR exceeded the total time bound before all pages were processed.")
                break
            try:
                recognized = subprocess.run([tesseract, str(image), "stdout", "--psm", "6", "tsv"], cwd=temporary, env=command_environment(), stdin=subprocess.DEVNULL, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, text=True, timeout=min(10, remaining), check=False)
            except (OSError, subprocess.TimeoutExpired):
                limitations.append(f"OCR page {page_number} recognition failed or timed out.")
                continue
            text, confidence = parse_tesseract_tsv(recognized.stdout if recognized.returncode == 0 else "")
            pages[page_number] = {"text": text[:10_000], "hasTextLayer": False, "extraction": "ocr", "confidence": confidence, "source": {"path": path.relative_to(workspace_root()).as_posix(), "page": page_number, "extraction": "ocr", "confidence": confidence}}
            if not text or confidence is None:
                limitations.append(f"OCR page {page_number} returned no text or confidence.")
            elif confidence < MIN_OCR_CONFIDENCE:
                limitations.append(f"OCR page {page_number} confidence is below the delivery threshold.")
    status = "passed" if pages and all(page_number in pages and pages[page_number].get("confidence") is not None and pages[page_number].get("confidence", 0) >= MIN_OCR_CONFIDENCE for page_number in page_numbers) else "unknown"
    return pages, {"status": status, "engine": "tesseract", "requestedPages": page_numbers, "processedPages": sorted(pages), "limitations": limitations}, limitations


def parse_tesseract_tsv(value: str) -> tuple[str, float | None]:
    words: list[str] = []
    confidences: list[float] = []
    for line in value.splitlines()[1:]:
        columns = line.split("\t")
        if len(columns) < 12:
            continue
        word = columns[11].strip()
        try:
            confidence = float(columns[10])
        except ValueError:
            continue
        if word and confidence >= 0:
            words.append(word)
            confidences.append(confidence / 100.0)
    return " ".join(words), sum(confidences) / len(confidences) if confidences else None


def command_environment() -> dict[str, str]:
    allowed = {"PATH", "LANG", "LC_ALL", "TMPDIR"}
    environment = {key: value for key, value in os.environ.items() if key in allowed}
    environment.setdefault("PATH", os.defpath)
    return environment


def xlsx_selector(request: dict[str, Any]) -> str | None:
    cell = request.get("cell")
    range_name = request.get("range")
    if (cell is None) == (range_name is None):
        return None
    selector = cell if isinstance(cell, str) else range_name
    if not isinstance(selector, str) or not selector or ":" in selector and selector.count(":") != 1:
        return None
    try:
        from openpyxl.utils.cell import range_boundaries
        min_col, min_row, max_col, max_row = range_boundaries(selector)
    except ValueError:
        return None
    cell_count = (max_col - min_col + 1) * (max_row - min_row + 1)
    if min_col < 1 or min_row < 1 or cell_count > MAX_XLSX_CELLS:
        return None
    return selector


def xlsx_cells(selector: str) -> list[str]:
    from openpyxl.utils.cell import get_column_letter, range_boundaries
    min_col, min_row, max_col, max_row = range_boundaries(selector)
    return [f"{get_column_letter(col)}{row}" for row in range(min_row, max_row + 1) for col in range(min_col, max_col + 1)]


def xlsx_values(worksheet: Any, selector: str) -> Any:
    from openpyxl.utils.cell import range_boundaries
    min_col, min_row, max_col, max_row = range_boundaries(selector)
    values = [[worksheet.cell(row=row, column=column).value for column in range(min_col, max_col + 1)] for row in range(min_row, max_row + 1)]
    if min_col == max_col and min_row == max_row:
        return values[0][0]
    return values


def xlsx_write_values(worksheet: Any, selector: str, request: dict[str, Any]) -> Any:
    from openpyxl.utils.cell import range_boundaries
    min_col, min_row, max_col, max_row = range_boundaries(selector)
    if min_col == max_col and min_row == max_row:
        value = request.get("value")
        if isinstance(value, (dict, list)):
            raise ValueError("cell_value_invalid")
        worksheet.cell(row=min_row, column=min_col).value = value
        return value
    values = request.get("values")
    if not isinstance(values, list) or len(values) != max_row - min_row + 1 or any(not isinstance(row, list) or len(row) != max_col - min_col + 1 for row in values):
        raise ValueError("range_values_invalid")
    for row_offset, row in enumerate(values):
        for col_offset, value in enumerate(row):
            if isinstance(value, (dict, list)):
                raise ValueError("cell_value_invalid")
            worksheet.cell(row=min_row + row_offset, column=min_col + col_offset).value = value
    return values


def xlsx_snapshot(workbook: Any) -> dict[str, Any] | None:
    snapshot: dict[str, Any] = {}
    count = 0
    for sheet in workbook.worksheets:
        if sheet.max_row * sheet.max_column > MAX_XLSX_CELLS:
            return None
        for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=1, max_col=sheet.max_column):
            for cell in row:
                snapshot[f"{sheet.title}!{cell.coordinate}"] = cell.value
                count += 1
                if count > MAX_XLSX_CELLS:
                    return None
    return snapshot


def changed_outside_target(before: dict[str, Any], after: dict[str, Any], sheet: str, selector: str) -> list[str]:
    target = {f"{sheet}!{cell}" for cell in xlsx_cells(selector)}
    changed: list[str] = []
    for key in set(before) | set(after):
        if key not in target and before.get(key) != after.get(key):
            changed.append(key)
    return changed


def write_minimal_docx(path: Path, paragraphs: list[str]) -> None:
    document = "".join(f"<w:p><w:r><w:t xml:space=\"preserve\">{escape(text)}</w:t></w:r></w:p>" for text in paragraphs)
    content_types = """<?xml version=\"1.0\" encoding=\"UTF-8\" standalone=\"yes\"?><Types xmlns=\"http://schemas.openxmlformats.org/package/2006/content-types\"><Default Extension=\"rels\" ContentType=\"application/vnd.openxmlformats-package.relationships+xml\"/><Default Extension=\"xml\" ContentType=\"application/xml\"/><Override PartName=\"/word/document.xml\" ContentType=\"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml\"/></Types>"""
    rels = """<?xml version=\"1.0\" encoding=\"UTF-8\" standalone=\"yes\"?><Relationships xmlns=\"http://schemas.openxmlformats.org/package/2006/relationships\"><Relationship Id=\"rId1\" Type=\"http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument\" Target=\"word/document.xml\"/></Relationships>"""
    document_xml = f"""<?xml version=\"1.0\" encoding=\"UTF-8\" standalone=\"yes\"?><w:document xmlns:w=\"http://schemas.openxmlformats.org/wordprocessingml/2006/main\"><w:body>{document}<w:sectPr/></w:body></w:document>"""
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", rels)
        archive.writestr("word/document.xml", document_xml)


def main() -> int:
    for line in sys.stdin:
        if not line.strip():
            continue
        try:
            request = json.loads(line)
            if not isinstance(request, dict):
                raise ValueError("request_must_be_object")
            payload = handle(request)
        except Exception:
            payload = response("unknown", ok=False, error="request_invalid")
        sys.stdout.write(json.dumps(payload, ensure_ascii=False, default=json_default) + "\n")
        sys.stdout.flush()
    return 0


def json_default(value: Any) -> str:
    """Keep valid XLSX date/time values from crashing the JSONL process."""
    isoformat = getattr(value, "isoformat", None)
    if callable(isoformat):
        return str(isoformat())
    return str(value)


if __name__ == "__main__":
    raise SystemExit(main())
