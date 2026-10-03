import os
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from pypdf import PdfWriter

import sys
sys.path.insert(0, str(Path(__file__).parents[1]))
import worker


class WorkerTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="helm-worker-test-")
        self.old_root = os.environ.get("HELM_WORKSPACE_ROOT")
        os.environ["HELM_WORKSPACE_ROOT"] = self.temp.name

    def tearDown(self):
        if self.old_root is None:
            os.environ.pop("HELM_WORKSPACE_ROOT", None)
        else:
            os.environ["HELM_WORKSPACE_ROOT"] = self.old_root
        self.temp.cleanup()

    def test_docx_artifact_and_inspect_are_bounded(self):
        created = worker.handle({"id": "d", "operation": "docx_create", "path": "report.docx", "paragraphs": ["Hello"], "runId": "run-1"})
        self.assertTrue(created["ok"])
        self.assertEqual(created["receipt"]["artifact"]["sourceRunId"], "run-1")
        self.assertTrue(zipfile.is_zipfile(Path(self.temp.name) / "report.docx"))
        inspected = worker.handle({"id": "i", "operation": "inspect", "path": "report.docx"})
        self.assertEqual(inspected["result"]["path"], "report.docx")

    def test_xlsx_range_write_and_read(self):
        path = Path(self.temp.name) / "book.xlsx"
        workbook = Workbook()
        workbook.active.title = "Data"
        workbook.save(path)
        written = worker.handle({"id": "w", "operation": "xlsx_write_range", "path": "book.xlsx", "sheet": "Data", "cell": "B2", "value": "ok"})
        self.assertTrue(written["ok"])
        read = worker.handle({"id": "r", "operation": "xlsx_read_range", "path": "book.xlsx", "sheet": "Data", "cell": "B2"})
        self.assertEqual(read["result"]["value"], "ok")
        self.assertEqual(load_workbook(path, data_only=True)["Data"]["B2"].value, "ok")

    def test_path_escape_and_unknown_pdf_text_are_explicit(self):
        escaped = worker.handle({"id": "x", "operation": "inspect", "path": "../secret.txt"})
        self.assertFalse(escaped["ok"])
        self.assertEqual(escaped["error"], "path_outside_workspace")
        pdf = Path(self.temp.name) / "blank.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        with pdf.open("wb") as stream:
            writer.write(stream)
        result = worker.handle({"id": "p", "operation": "pdf_extract", "path": "blank.pdf"})
        self.assertFalse(result["ok"])
        self.assertEqual(result["error"], "unknown_text_layer")


if __name__ == "__main__":
    unittest.main()
